from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/Stakeholder-Year-End-Outlook.docx")
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(89, 89, 89)


def font(run, size=10.5, bold=False, color=BLACK):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def body(doc, text, lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    if lead:
        r = p.add_run(lead)
        font(r, bold=True)
    r = p.add_run(text)
    font(r)
    return p


def heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(15)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    font(r, size=14, bold=True)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    r = p.add_run(text)
    font(r)
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    style = doc.styles["Normal"]
    style.font.name = "Aptos"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    style.font.size = Pt(10.5)
    for name in ("Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = BLACK
        doc.styles[name]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        doc.styles[name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(6)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = title.add_run("Year End Outlook for Therapist Sidekick")
    font(r, size=24, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    r = subtitle.add_run("What an accelerated AI timeline can credibly accomplish by 31 December 2026")
    font(r, size=11, color=GRAY)

    body(doc, "By year end, the most credible goal is a limited design-partner pilot of a therapist session-review workflow. It should allow a therapist to move from a recorded session to a reviewable transcript, synchronized audio, and a clearly labeled AI draft summary. The purpose is to test whether this reduces administrative burden while preserving therapist judgment and control.", "Recommendation. ")
    body(doc, "This is a meaningful outcome, but it is not a finished practice-management platform, a broad client product, or autonomous clinical AI. Keeping the commitment narrow will make it faster to build and safer to evaluate.")

    heading(doc, "Starting point")
    body(doc, "The project already has an early working prototype. It records from the browser, sends audio to Amazon HealthScribe, stores timestamped and speaker-labeled transcript segments, plays the matching audio, lets a therapist revise speaker names, and displays a draft clinical note. The AWS infrastructure also provides private, encrypted storage for the transcription flow.")
    body(doc, "The main work before a real pilot is not simply improving the model. The product needs secure therapist access, clear data boundaries, consent and retention controls, auditability, recovery from processing failures, a usable client and session organization model, and direct feedback from clinicians.")

    heading(doc, "What the pilot should include")
    body(doc, "The pilot should focus on the session workflow and make every AI output reviewable. A small group of named design partners should be able to sign in, work only with their authorized data, and use a recorded consent step before a session is captured or processed.")
    bullet(doc, "A dependable capture or upload process with visible processing status, retry handling, and synchronized transcript-to-audio review.")
    bullet(doc, "Transcript and speaker-label correction, so therapists can finalize or improve the source material before relying on it.")
    bullet(doc, "A draft summary and key-points view that is editable, optional, clearly not a clinical conclusion, and traceable back to the source transcript where practical.")
    bullet(doc, "Simple therapist-controlled organization of clients and sessions, so pilot users are not navigating local folders or technical identifiers.")
    bullet(doc, "Structured feedback and quality measurement for transcript errors, summary faithfulness, review time, usability, and operational failures.")

    heading(doc, "What should remain out of scope")
    body(doc, "The team should not promise a full practice-management suite by December. Scheduling, billing, insurance, multi-provider administration, and client discovery are substantial products in their own right. A broad iPhone client experience also deserves separate safety, permissions, and product work.")
    body(doc, "The project should also avoid promising a custom model that replaces HealthScribe this year. A responsible replacement requires governed data, representative evaluation, measurable baselines, monitoring, and proof that it is materially better for therapy sessions. Automated risk detection, diagnosis, and claims about emotion or tone should remain outside the pilot's scope.")

    heading(doc, "Suggested delivery path")
    body(doc, "September should establish the pilot boundary: which therapists will participate, what the intended use is, how consent works, how long data is retained, and how success will be measured. October should harden the core workflow with identity, access control, client/session organization, audit events, and resilient processing. November should add therapist-controlled draft synthesis and run structured quality reviews. December should operate the small pilot, resolve high-priority issues, and produce evidence for a broader-beta decision.")

    heading(doc, "What accelerated AI changes")
    body(doc, "AI-assisted development can materially speed up interface work, integration, testing, transcript normalization, retrieval prototypes, and internal evaluation tooling. It can also accelerate iteration on the format of therapist-facing summaries. That is why a focused pilot is plausible within the remaining months.")
    body(doc, "It does not remove the need for clinician trust, privacy and consent decisions, security controls, representative evaluation, or safe product design. In this domain, those are release gates rather than tasks that can be compressed away.")

    heading(doc, "Conditions for success")
    bullet(doc, "A dedicated small team, roughly three to five full-time product and engineering contributors, with rapid decisions and clear ownership.")
    bullet(doc, "A clinical lead and privacy/security/legal support that can define the pilot's consent, retention, intended-use, and review policies before real sessions are processed.")
    bullet(doc, "Three to five committed design partners who give recurring feedback and agree to use the controlled pilot workflow.")
    bullet(doc, "Continued use of HealthScribe as the baseline while the team gathers evidence, instead of diverting capacity into a premature model-replacement program.")

    heading(doc, "How to judge the December result")
    body(doc, "Before the pilot starts, stakeholders should agree on specific thresholds. The review should cover: whether consent and access controls worked as designed; how often sessions completed or required recovery; how much transcript correction was needed; whether summaries were accepted, edited, or rejected; whether therapists could verify material content against source; and whether design partners would choose to keep using the workflow.")
    body(doc, "The appropriate December decision is evidence-based. Expand into a broader beta only if the pilot demonstrates clinician value, dependable operation, and the agreed safety and control measures. If it does not, the outcome is still useful: it will identify the specific reliability, trust, or workflow gaps that must be addressed next.")

    heading(doc, "Decision requested")
    body(doc, "Approve the narrow session-review pilot as the year-end objective. Assign the clinical, privacy, security, product, engineering, and design-partner capacity needed to support it. Defer full practice operations, broad client features, and custom-model claims until the pilot produces evidence that the core workflow is valuable and trustworthy.")

    doc.core_properties.title = "Year End Outlook for Therapist Sidekick"
    doc.core_properties.subject = "Stakeholder outlook for an accelerated AI timeline"
    doc.core_properties.author = "Therapist Sidekick"
    doc.save(OUT)


if __name__ == "__main__":
    main()
