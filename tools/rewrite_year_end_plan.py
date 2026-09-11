from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/Stakeholder-Year-End-Outlook.docx")
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(88, 88, 88)


def set_font(run, size=11, bold=False, color=BLACK, italic=False):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def paragraph(doc, text, lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.18
    if lead:
        r = p.add_run(lead)
        set_font(r, bold=True)
    r = p.add_run(text)
    set_font(r)
    return p


def heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(17)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    set_font(r, size=14, bold=True)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.17)
    r = p.add_run(text)
    set_font(r)
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.88)
    section.right_margin = Inches(0.88)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(11)
    for style_name in ("Heading 1", "Heading 2", "Title"):
        style = doc.styles[style_name]
        style.font.color.rgb = BLACK
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(6)
    r = title.add_run("End of Year Product Demonstration Plan")
    set_font(r, size=24, bold=True)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(18)
    r = sub.add_run("Therapist Sidekick | September through the end of 2026")
    set_font(r, size=11, color=GRAY)

    paragraph(doc, "The goal for the rest of this year is to make Therapist Sidekick feel real to a therapist. By the end of the year, I want to be able to sit with a therapist, let them click through the product, and give them a clear sense of how it could fit into an ordinary practice. They should be able to use it with sample information and, where the right consent and safeguards are in place, try parts of the flow with a real client. The point is not to launch a finished product or accept broad public use. It is to create a convincing, usable version of the experience that can be shown, tested, and improved.")
    paragraph(doc, "This is a solo project, so the plan needs to be practical. I am not trying to build every part of a therapy platform by the end of the year. I am trying to build the parts that make the overall idea understandable: a therapist-facing workspace, a simple client-facing experience, a believable session flow, and a foundation for handling sensitive information with care.")

    heading(doc, "What I want a therapist to experience")
    paragraph(doc, "A therapist should be able to open the web application and feel as though it is a calm place to manage the day-to-day shape of a practice. They should be able to add or view clients, see a simple calendar, set up a recurring appointment pattern, and keep basic billing information organized. At this stage, the system does not need to charge a credit card, submit an insurance claim, or become the official calendar of record. It does need to make those future steps feel believable and show how they would fit into the rest of the product.")
    paragraph(doc, "The session flow is where the product should become especially tangible. A therapist should be able to select a client, begin or upload a session recording, and later return to a session page that shows the audio beside the written conversation. Clicking a part of the transcript should take the listener to the relevant point in the recording. If speakers are labeled incorrectly, the therapist should be able to correct them. A summary or note can be present if it is helpful, but it should be clearly treated as a draft rather than an answer the therapist must trust.")
    paragraph(doc, "The experience does not need to be perfect to be valuable. It needs to work smoothly enough that a therapist can imagine using it after a real session. That means the important details matter: the right client should be easy to find, the session should feel connected to the calendar, the recording should not feel lost after it is processed, and the transcript should feel like something a therapist can review rather than a technical artifact.")

    heading(doc, "What I want a client to experience")
    paragraph(doc, "I also want a simple client-facing web application that makes the other side of the relationship understandable. A test user should be able to browse or select a therapist, see available or proposed appointment times, manage a calendar, and see a small history of past sessions. The client experience should be much simpler than the therapist workspace. It should feel supportive and easy to follow, not like an electronic medical record.")
    paragraph(doc, "For a client who has already had a few sessions, the product should be able to show the session materials that the therapist has chosen to make available. That could include the audio and transcript together, with clear controls and a straightforward explanation of what is being shown. Whether a client can see a particular recording or transcript should always be controlled by the therapist and by the permissions built into the product. For this first version, the main goal is to make the shared journey visible, not to settle every possible client-access policy.")

    heading(doc, "What is in scope by the end of the year")
    paragraph(doc, "The end-of-year version should focus on a coherent demonstration rather than a large list of disconnected features. I would consider the following a strong outcome:")
    bullet(doc, "A therapist web workspace with clients, session history, a simple calendar, recurring appointment patterns, and a place to organize billing status or notes without actually processing payments.")
    bullet(doc, "A client web experience where a test user can select or view a therapist, manage a simple appointment calendar, and see the sessions and materials the therapist has permitted them to view.")
    bullet(doc, "A session-review experience that brings together the recording, transcript, speaker labels, timestamps, and basic draft notes in one understandable place.")
    bullet(doc, "A storage approach that keeps users, clients, appointments, sessions, recordings, transcripts, and permissions connected in a clear way. It does not need to be the final architecture, but it should be clean enough that I can keep building on it with confidence rather than replacing it immediately.")
    bullet(doc, "A clear way to show what is real today, what is simulated for the demonstration, and what will need more work before the product is used in routine practice.")

    heading(doc, "Therapist roles and permissions")
    paragraph(doc, "The first version should also make it clear that a practice may involve more than one kind of therapist. A main therapist should be able to manage their own clients, appointments, sessions, and documentation. A supervisory therapist should be able to look at the information that has been shared with them and offer guidance, while not automatically having the same ability to change or finalize the client's record. A fill-in therapist should be able to cover a client for a limited period when that access has been explicitly granted. These distinctions should be understandable in the product itself, not just hidden in the code.")
    paragraph(doc, "The user interface does not need to cover every possible practice structure yet. It should give a therapist or practice owner a straightforward way to invite another clinician, choose the role they need, decide which clients or areas of the system they can access, and remove that access when it is no longer needed. A client should only see the therapists and session materials that have been made available to them.")
    paragraph(doc, "The storage design should support this from the beginning. Instead of assuming that one therapist owns every client forever, the application should store the relationship between a person, their role, the practice or account they belong to, and the scope of access they have. For example, access may be limited to a particular client, a group of clients, scheduling only, session review only, or read-only supervision. Every important change to these permissions should be recorded so that it is possible to understand who had access and why. This will make the demonstration more realistic now and prevent the underlying design from becoming a barrier as the product grows.")

    heading(doc, "What I am deliberately not trying to finish")
    paragraph(doc, "I am not trying to make Therapist Sidekick ready for widespread real-world use by the end of the year. I am also not trying to build a full billing system, insurance workflow, complete marketplace, mobile application, multi-practice administration system, or a custom clinical AI model. Those things may become important later, but they would pull attention away from the experience I need to prove first.")
    paragraph(doc, "The current transcription, diarization, and draft-note tools are useful because they make the session flow possible now. They do not have to be the final answer. Exploring post-training or adapting a model for therapy is an alternate objective that I am very interested in pursuing alongside the main product work. It is not the main priority for the end of the year. It should move forward as a careful research path, especially once there are trusted therapist relationships, a clear permission model, and a meaningful way to measure whether a new model is actually better for the people using it.")

    heading(doc, "A practical path through the end of the year")
    paragraph(doc, "The first part of the work is to make the core structure feel solid. I would connect the existing recording and transcript work to a durable model of therapists, clients, appointments, sessions, and access permissions. At the same time, I would build the basic therapist and client pages so that the product starts to feel like one connected system instead of a recorder sitting next to a set of future ideas.")
    paragraph(doc, "The next part is to fill in the flows that make the demonstration believable. For the therapist, that means client organization, calendar views, recurring appointments, billing organization, and the session-review page. For the client, that means a simpler view of therapist selection, appointments, and permitted session history. A recurring schedule is a reasonable feature to include because it makes a therapist's calendar feel much closer to a real practice even before there is a complete scheduling engine.")
    paragraph(doc, "The final part is rehearsal. I should be able to create a few realistic test accounts and walk through the product from both sides: a therapist organizes a client, establishes a recurring appointment, records or uploads a session, reviews the transcript and audio, and then decides what the client can see. The client can sign in, see the appointment history, and review the materials that have been made available. If this can be demonstrated cleanly from beginning to end, it will provide a much stronger basis for conversations with therapists than an isolated feature demo.")

    heading(doc, "Storage and technical foundations")
    paragraph(doc, "The storage design should be good enough to support the product as it grows, while staying simple enough to build alone. The important idea is that each record has a clear home and relationship to the others. A therapist belongs to an account. A client is connected to the therapist or practice that is allowed to work with them. An appointment belongs to that therapist and client. A session belongs to the appointment or client relationship. The recording, transcript, speaker corrections, and draft notes belong to the session. Access to client-facing materials is a separate, explicit choice rather than an accidental side effect of storing the file.")
    paragraph(doc, "The current AWS setup already gives the project a useful starting point: private storage buckets, encryption at rest, and access roles for the transcription workflow. The next step is to move away from relying on local folders as the organizing system and toward an application database that records what exists, who can see it, how long it should be retained, and what should happen when it is deleted. This will not make the architecture final, but it will make it much easier to refine safely over time.")

    heading(doc, "Privacy, HIPAA, and recording law")
    paragraph(doc, "This project should not present itself as HIPAA compliant or legally ready for routine use before that work has been completed. Building a clear plan, documenting the design choices, and identifying the questions that need professional review are important supporting objectives that I will pursue alongside the main product work. They are not a reason to delay the demonstration experience, but they are necessary for making responsible decisions about when and how real information can be used.")
    paragraph(doc, "The first step is to decide exactly when real information may enter the system. Sample or synthetic data should be the default for demonstrations. If a real session is ever used for testing, there should be a written process for informed consent, a clear explanation of what will be recorded and stored, a way to stop or withdraw from the process where appropriate, and an explicit decision about whether the client can access the material. The consent language and workflow should be reviewed by qualified counsel and by the therapist who is involved.")
    paragraph(doc, "HIPAA depends on the role the product will play and the way it is used by a therapist or practice. As a parallel effort, I will research the product's role, identify every vendor that creates, receives, maintains, or transmits protected health information, and learn which business associate agreements may be required. Before supporting routine use with protected health information, these questions should be confirmed with healthcare privacy counsel. The current design should also be checked against a written risk analysis covering access controls, encryption, backups, audit records, incident response, retention, deletion, and the ability to give records back to the appropriate party. AWS and any transcription or hosting providers should be evaluated based on the actual services and contractual terms that will be used, not on a general claim that a cloud platform is automatically compliant.")
    paragraph(doc, "Recording law is another side objective that deserves serious attention as the product develops. Consent requirements can vary by state and may depend on where each participant is located, whether the session is remote, and the way recording is disclosed. Federal law is not the only rule that matters. Before recording real sessions beyond a carefully reviewed pilot, I will seek a jurisdiction-by-jurisdiction review, a practical rule for remote sessions involving different states, and consent language that is easy for both therapist and client to understand. This document is a product plan, not legal advice; the final policy and user-facing language should come from qualified legal counsel.")

    heading(doc, "What success looks like at the end of the year")
    paragraph(doc, "Success is being able to put the product in front of a therapist and have the conversation move beyond the abstract. The therapist should be able to click through a believable practice workflow, understand what would happen before and after a session, see how a client would experience their side of the system, and tell me where the product is useful or uncomfortable. I should leave that conversation with feedback about real decisions rather than guesses about whether the general idea is interesting.")
    paragraph(doc, "A second form of success is having a foundation I am comfortable extending. The product should have a coherent set of records, sensible permissions, private storage, documented assumptions, and a clear next list of compliance and legal questions. It does not need to be ready to invite the public in. It needs to be good enough that the next phase can be guided by therapist feedback, not by a fragile prototype or an unclear architecture.")

    heading(doc, "Immediate next steps")
    paragraph(doc, "The practical next step is to turn this document into a small, ordered build list: first the accounts, therapist roles, permissions, and storage model; then therapist client and calendar management; then the client-facing pages; then the end-to-end session demonstration; and finally the documentation and review needed before involving real session content. In parallel, I should begin conversations with a therapist and pursue privacy and recording-law research so that the demonstration is built around the right questions from the beginning.")

    heading(doc, "Main objectives at a glance")
    bullet(doc, "Make a therapist-facing product that feels real enough to demonstrate a normal practice flow from client setup and recurring appointments through session review.")
    bullet(doc, "Make a simple client-facing web experience where a test user can select a therapist, manage an appointment calendar, and view the session materials that have been explicitly shared.")
    bullet(doc, "Create a storage and permissions foundation that connects accounts, therapists, supervisory therapists, fill-in therapists, clients, appointments, sessions, recordings, transcripts, and access decisions in a clear way.")
    bullet(doc, "Give main therapists, supervisory therapists, and fill-in therapists an understandable way to manage their access through the user interface, with clear limits on who can view, advise on, or change client documentation.")
    bullet(doc, "Keep the transcription and draft-note flow useful for the demonstration without making custom clinical AI the main end-of-year commitment.")
    bullet(doc, "Pursue legal and privacy research alongside the product work, so that the next phase has a clear path for handling consent, recordings, protected health information, and future routine use.")

    doc.core_properties.title = "End of Year Product Demonstration Plan"
    doc.core_properties.subject = "Guidelines for Therapist Sidekick through the end of 2026"
    doc.core_properties.author = "Therapist Sidekick"
    doc.save(OUT)


if __name__ == "__main__":
    main()
