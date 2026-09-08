from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/Stakeholder-Year-End-Outlook.docx")
NAVY = RGBColor(24, 52, 76)
TEAL = RGBColor(13, 110, 113)
INK = RGBColor(35, 43, 53)
MUTED = RGBColor(91, 102, 114)
PALE_BLUE = "EAF3F7"
PALE_TEAL = "E7F3F1"
PALE_GOLD = "FFF4DB"
WHITE = "FFFFFF"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def set_cell_text(cell, value, bold=False, color=INK, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(value)
    run.bold = bold
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.22 + level * 0.18)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.font.color.rgb = INK
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Aptos Display"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    r.font.size = Pt(16 if level == 1 else 12)
    r.font.color.rgb = NAVY if level == 1 else TEAL
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        r.font.color.rgb = INK
    r = p.add_run(text)
    for run in p.runs:
        run.font.name = "Aptos"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        run.font.size = Pt(10.5)
        run.font.color.rgb = INK
    return p


def add_callout(doc, title, body, fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.7)
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell, 140, 180, 140, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = NAVY
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    r2.font.size = Pt(10)
    r2.font.color.rgb = INK
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = Inches(widths[idx])
        shade(cell, NAVY.__str__().replace("'", "") if False else "18344C")
        set_cell_text(cell, text, bold=True, color=RGBColor(255, 255, 255), size=9)
    set_repeat_table_header(table.rows[0])
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cells[idx].width = Inches(widths[idx])
            shade(cells[idx], WHITE if row_idx % 2 == 0 else "F5F8FA")
            set_cell_text(cells[idx], text, size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Therapist Sidekick | Stakeholder discussion draft | 8 September 2026")
    r.font.name = "Aptos"
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)
    footer(sec)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("THERAPIST SIDEKICK")
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = TEAL

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Year-end outlook: what an accelerated AI timeline can realistically deliver")
    r.bold = True
    r.font.name = "Aptos Display"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    r.font.size = Pt(23)
    r.font.color.rgb = NAVY

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("Stakeholder discussion draft | 8 September 2026 | Horizon: 31 December 2026")
    r.font.size = Pt(10.5)
    r.font.color.rgb = MUTED

    add_callout(doc, "Bottom line", "With a focused, accelerated effort, the credible year-end outcome is a secure, consented design-partner pilot of the therapist session-review workflow - not a general-purpose production practice-management platform or autonomous clinical AI. The value to prove is that a therapist can move from recorded session to reviewable transcript, source-linked draft synthesis, and useful session context with materially less administrative burden.", PALE_TEAL)

    add_heading(doc, "What exists today")
    add_body(doc, "The repository contains a working local prototype: browser microphone capture, asynchronous Amazon HealthScribe processing, timestamped and speaker-labeled transcript segments, synchronized audio playback, editable speaker labels, draft clinical-note display, and encrypted AWS storage infrastructure. The web portal is an early single-workspace experience.")
    add_body(doc, "The gap to a real pilot is not principally model capability. It is product hardening: identity and role-based access, consent and retention controls, organization/client data boundaries, auditability, failure handling, usability research, clinical policy, and operating discipline.")

    add_heading(doc, "Recommended December commitment")
    add_body(doc, "Deliver a tightly scoped 'Therapist Session Review Pilot' for a small number of design partners in a controlled environment. It should validate the first and highest-priority product loop: capture, review, correct, synthesize, and learn from a therapy session while keeping clinical judgment with the therapist.")

    add_table(doc,
        ["Included in the pilot", "Evidence of completion"],
        [
            ("Secure therapist access and tenant separation", "Named pilot users can sign in; access is limited to their authorized data; key access and operational events are auditable."),
            ("Consent-aware session workflow", "Recording and processing occur only after an explicit, recorded consent step; retention and deletion rules are visible and testable."),
            ("Reliable capture-to-review path", "A therapist can upload or record a session, follow processing status, play synchronized audio, inspect timestamps, and correct speaker labels/transcript content."),
            ("Reviewable AI draft synthesis", "A therapist receives a clearly labeled draft summary/key-points view, can edit or reject it, and can trace material points back to source transcript passages."),
            ("Therapist-controlled client/session organization", "Pilot users can find the right client and session, attach reviewed artifacts, and return to them without relying on local folders or opaque IDs."),
            ("Pilot learning and quality loop", "The team captures structured feedback and a small, governed evaluation set to measure transcript quality, summary faithfulness, review friction, and failures."),
        ], [2.25, 4.45])

    add_heading(doc, "What should not be promised by year end")
    add_table(doc,
        ["Do not frame as year-end delivery", "Why it should remain later-stage"],
        [
            ("Full practice-management suite", "Scheduling, billing, insurance, multi-provider administration, and client discovery each create separate workflow, compliance, and integration commitments."),
            ("Broad client iPhone product", "A meaningful client experience needs its own permissions model, safety design, content boundaries, and sustained product work."),
            ("Custom foundation models or replacement of HealthScribe", "A responsible replacement needs governed data, representative evaluation, baselines, monitoring, and proof of material advantage - none should be compressed into a few months."),
            ("Automated clinical conclusions, risk detection, or emotion/tone judgments", "These are explicitly outside the product's near-term safety boundary and require validation that does not fit an accelerated pilot."),
            ("Large-practice / enterprise rollout", "Organization controls, support, contracting, security review, and operational maturity must follow evidence from the core workflow."),
        ], [2.45, 4.25])

    add_heading(doc, "Accelerated path from September through December")
    add_table(doc,
        ["Period", "Primary objective", "Decision-quality output"],
        [
            ("September", "Pilot definition and foundation", "Choose 3-5 design partners; define intended use, consent, retention, and non-use boundaries; map the current prototype's gaps; establish quality and safety metrics."),
            ("October", "Pilot-ready core workflow", "Implement secure identity/data boundaries, consent flow, session/client organization, durable storage and audit events; harden upload, status, retry, and review UX."),
            ("November", "Clinician-controlled AI and measurement", "Release editable draft synthesis with source references; add transcript correction/finalization; run structured usability and quality reviews against the HealthScribe baseline."),
            ("December", "Controlled design-partner pilot", "Operate the workflow with a deliberately limited cohort; triage issues rapidly; deliver a year-end evidence pack and a go/no-go recommendation for broader beta."),
        ], [1.05, 2.55, 3.1])

    add_heading(doc, "What 'accelerated AI' changes - and what it does not")
    add_body(doc, "AI coding tools can substantially compress interface construction, integration work, test generation, transcript normalization, retrieval prototypes, and internal evaluation tooling. They can also let the team iterate faster on therapist-facing summary formats. That makes the pilot scope above plausible with a small, capable team.")
    add_body(doc, "AI does not eliminate the parts that make this domain difficult: consent, privacy, security, clinical policy, safe interaction design, representative evaluation, and clinician trust. Those are product gates, not implementation backlog. The fastest path is therefore narrower and more measured, not broader.")

    add_heading(doc, "Conditions required for this outcome")
    add_bullet(doc, "A dedicated delivery team of roughly 3-5 full-time product/engineering contributors, with accountable product leadership and rapid decision-making.")
    add_bullet(doc, "A named clinical lead plus privacy/security/legal support empowered to set intended-use, consent, retention, and review policies before real session data is processed.")
    add_bullet(doc, "Three to five committed design partners who can provide recurring feedback and use only the approved pilot workflow.")
    add_bullet(doc, "A deliberate choice to use HealthScribe as the early baseline and to avoid a model-replacement program until quality evidence justifies it.")
    add_bullet(doc, "A controlled pilot environment, not an implied HIPAA/enterprise-ready commercial launch, unless the required compliance evidence is separately established.")

    add_heading(doc, "Year-end success measures")
    add_body(doc, "Set exact targets before the pilot starts. The measures below are the right decision frame; they should be reported with sample size, cohort, and known limitations rather than treated as marketing claims.")
    add_table(doc,
        ["Dimension", "Evidence to collect"],
        [
            ("Safety and control", "Consent coverage; authorization/audit checks; retention/deletion test results; count and resolution of privacy or access incidents."),
            ("Workflow reliability", "Sessions started, completed, failed, and recovered; processing latency; uptime/defect trends; success through the complete review path."),
            ("Transcript usefulness", "Therapist review time; correction patterns; speaker-label quality; timestamp-to-audio alignment; representative error taxonomy."),
            ("Synthesis trust", "Therapist accept/edit/reject behavior; source-traceability coverage; faithfulness review; qualitative reasons the draft was or was not useful."),
            ("User value", "Design-partner willingness to continue; reported administrative burden reduction; repeat use; high-priority workflow improvements for the next phase."),
        ], [1.7, 5.0])

    add_heading(doc, "Stakeholder decisions needed now")
    add_bullet(doc, "Approve the narrow year-end pilot outcome as the top company objective, and explicitly defer full practice operations, broad client features, and custom-model claims.")
    add_bullet(doc, "Fund or assign the delivery, clinical, security, and design-partner capacity required by the conditions above.")
    add_bullet(doc, "Authorize a pilot governance package: intended use, consent language/process, data retention/deletion policy, access model, incident process, and feedback protocol.")
    add_bullet(doc, "Agree that December's decision is evidence-based: continue into a broader beta only if clinician value, operational reliability, and safety controls meet the agreed thresholds.")

    add_callout(doc, "Suggested stakeholder message", "By year end, we can credibly demonstrate a therapist-controlled, AI-assisted session-review workflow with selected design partners. Our aim is not to claim that AI can safely automate therapy or that we have finished a practice platform. Our aim is to establish the trustworthy core, measure whether it reduces therapist burden, and use that evidence to choose the next investment.", PALE_GOLD)

    doc.core_properties.title = "Therapist Sidekick Year-End Outlook"
    doc.core_properties.subject = "Accelerated AI delivery outlook for stakeholders"
    doc.core_properties.author = "Therapist Sidekick"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    main()
