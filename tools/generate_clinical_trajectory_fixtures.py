"""Generate fictional clinical trajectories, session plans, and Gemini TTS voice tests."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


INK = "24313A"
MUTED = "65737D"
TEAL = "377A78"
LIGHT_TEAL = "EAF3F2"
SAND = "F5F0E8"
WHITE = "FFFFFF"
GRID = "D9D9D9"


CASES = [
    {
        "name": "Jordan Kim",
        "folder": "Jordan Kim",
        "email": "jordan.kim@example.com",
        "age": 32,
        "identity": "Korean American; nonbinary; they/them",
        "life": "Product designer living alone and working on a fast-moving remote team.",
        "setting": "Outpatient individual therapy",
        "acuity": "Lower acuity",
        "voice": "Puck",
        "voice_style": "Warm mid-range voice; thoughtful pace; sentences accelerate when discussing work; restrained humor.",
        "speaker": "JORDAN",
        "therapist_voice_fixture": "therapist-calm-01",
        "reason": "Work pressure is disrupting sleep, recovery time, and confidence in everyday decisions.",
        "tension": "Rest feels irresponsible unless every task is complete, but the definition of complete keeps moving.",
        "strengths": "Curious, observant, collaborative, financially stable, connected to two close friends, and willing to test small changes.",
        "backstory": [
            "Jordan built a career by being quick, responsive, and unusually good at anticipating revisions. A recent promotion brought more autonomy but also ambiguous expectations and constant visibility across time zones.",
            "They began checking work messages in bed and waking before dawn with an internal list already running. Nothing dramatic has collapsed, yet ordinary choices now feel loaded and weekends have become recovery periods rather than a life of their own.",
        ],
        "themes": [
            "Performance-based self-worth and difficulty recognizing a sufficient stopping point.",
            "Sleep disruption maintained by late work, anticipatory checking, and self-criticism.",
            "Avoidance of direct expectation-setting with a manager.",
            "A wish to reclaim creativity and friendship without turning either into another goal.",
        ],
        "formulation": "Jordan's distress may be maintained by ambiguous workplace expectations, threat-focused checking, inconsistent recovery, and a belief that visible availability prevents disappointment.",
        "protective": "Stable housing, supportive friends, reflective capacity, meaningful creative interests, and no current indication of acute safety concerns in the fictional arc.",
        "relational": "Jordan is polite and analytically fluent, then becomes more spontaneous when they trust that uncertainty will not be graded.",
        "stance": "Collaborative behavioral observation, sleep-supportive routines, values clarification, and rehearsal of specific workplace conversations.",
        "safety": [
            "Do not inflate ordinary stress into a hidden disorder or crisis.",
            "Track sleep and functioning without presenting a generated score as a clinical conclusion.",
            "Keep improvement partial and believable; workload changes as well as personal habits matter.",
        ],
        "phases": [
            ("1", "1-2", "Map the work-sleep-checking loop and establish a baseline.", "Jordan notices that checking briefly reduces anxiety but extends the workday."),
            ("2", "3-4", "Test boundaries and tolerate uncertainty about others' reactions.", "They ask their manager to define response-time expectations."),
            ("3", "5-6", "Consolidate a sustainable stopping ritual and reconnect with non-work identity.", "A difficult week becomes a course correction rather than proof of failure."),
        ],
        "sessions": [
            ("The moving finish line", "Jordan describes a normal-looking life with no reliable off switch.", "They identify message checking as both relief and fuel for vigilance.", "Track two evenings without trying to change them."),
            ("The hour before sleep", "They map screens, unfinished tasks, and early waking without moralizing the routine.", "Jordan chooses one modest wind-down cue and refuses a complicated optimization plan.", "Compare sleep and urge-to-check notes."),
            ("What counts as responsive", "A manager's vague request triggers weekend work and resentment.", "Jordan rehearses a neutral question about deadlines and escalation channels.", "Send the clarification during work hours."),
            ("Silence after the boundary", "The manager responds briefly, and Jordan interprets the brevity as anger.", "They tolerate not sending an apologetic follow-up and examine alternative explanations.", "Notice the body response while waiting."),
            ("A weekend that belongs to someone", "Jordan accepts a friend's invitation but feels distracted and guilty.", "Pleasure and anxiety coexist without making the outing a failure.", "Choose one recurring non-work commitment."),
            ("Enough for this season", "They review gains, a recent late-night relapse, and remaining structural stress.", "Jordan names warning signs and a flexible reset plan rather than a perfect routine.", "Continue monitoring only when it is useful."),
        ],
        "dialogue": [
            ("THERAPIST", "When did work end yesterday, if it ended at all?"),
            ("JORDAN", "Officially? Six. In the version where I tell the truth, probably eleven-thirty."),
            ("THERAPIST", "What kept the day open?"),
            ("JORDAN", "A message I did not need to answer. I knew that. I answered it anyway, and then there was another one."),
            ("THERAPIST", "What did answering promise you in that moment?"),
            ("JORDAN", "That no one could be disappointed before I woke up. [pause] Which sounds impossible when I say it out loud."),
            ("THERAPIST", "Impossible, and still compelling."),
            ("JORDAN", "Exactly. I am not confused about the logic. I am trapped by how convincing it feels at ten forty-five."),
            ("THERAPIST", "Could we study that moment instead of asking you to win it every night?"),
            ("JORDAN", "That feels less like homework. Yes. I can write down what I think will happen if I leave the message alone."),
            ("THERAPIST", "And what actually happens by morning."),
            ("JORDAN", "Right. Data, but not a performance review. Please put that in bold."),
        ],
    },
    {
        "name": "Avery Chen",
        "folder": "Avery Chen",
        "email": "avery.chen@example.com",
        "age": 39,
        "identity": "Chinese American woman; she/her",
        "life": "Public-library branch manager sharing caregiving responsibilities for her father with two siblings.",
        "setting": "Outpatient individual therapy",
        "acuity": "Moderate",
        "voice": "Aoede",
        "voice_style": "Clear, grounded voice; measured delivery; becomes clipped when resentful and softer when naming grief.",
        "speaker": "AVERY",
        "therapist_voice_fixture": "therapist-calm-02",
        "reason": "Caregiving demands and unresolved family conflict have displaced sleep, friendships, and personal routines.",
        "tension": "Avery wants help but fears that asking directly will expose how angry and exhausted she already feels.",
        "strengths": "Practical, caring, organized, able to repair relationships, supported by a partner, and motivated to preserve connection with her father.",
        "backstory": [
            "Avery became the default organizer after her father's health declined. Her siblings contribute, but plans often reach Avery only after a problem has become urgent.",
            "She prides herself on handling logistics calmly. Over the past year, that competence has hardened into resentment, and she has begun avoiding calls because every ring feels like another assignment.",
        ],
        "themes": [
            "Caregiver strain, anticipatory grief, and anger that arrives before sadness.",
            "Uneven family labor and indirect communication about responsibility.",
            "Guilt when caregiving limits are interpreted as lack of love.",
            "Rebuilding ordinary routines that are not contingent on a crisis-free week.",
        ],
        "formulation": "Avery's strain may be reinforced by unclear family roles, emergency-driven planning, suppressed anger, and a belief that competent care requires constant availability.",
        "protective": "Supportive partner, stable work, caring siblings despite conflict, practical problem-solving skills, and willingness to have difficult conversations.",
        "relational": "Avery initially presents a detailed logistics report and redirects emotion into planning. Progress appears when she can express anger without immediately retracting it.",
        "stance": "Validate the real workload, distinguish care from sole responsibility, practice requests, and leave room for grief that cannot be solved operationally.",
        "safety": [
            "Do not portray caregiver frustration as evidence that Avery is unsafe or uncaring.",
            "Avoid assigning motives to siblings; scripts should preserve incomplete information and competing constraints.",
            "Any medical detail about the father remains minimal and fictional because the client arc is Avery's experience.",
        ],
        "phases": [
            ("1", "1-2", "Make the hidden workload visible and separate urgency from habit.", "Avery names resentment without converting it into a task list."),
            ("2", "3-5", "Practice concrete requests and tolerate imperfect help.", "One sibling takes a recurring responsibility in a different style than Avery would choose."),
            ("3", "6-8", "Approach anticipatory grief and build a sustainable care network.", "A setback no longer erases Avery's limits or her access to support."),
        ],
        "sessions": [
            ("The family switchboard", "Avery inventories the calls, appointments, and decisions that route through her.", "She distinguishes tasks only she can do from tasks she automatically absorbs.", "Bring one recent week of examples."),
            ("Anger before the phone rings", "A missed sibling callback triggers a disproportionate argument.", "Avery identifies anger as information about overload, not a verdict on the relationship.", "Draft one specific request."),
            ("The request without the case", "She rehearses asking a sibling to own medication pickup every Thursday.", "Avery notices how much evidence she wants to present before making a simple request.", "Send a brief request without overexplaining."),
            ("Help in someone else's style", "The sibling follows through but organizes the task differently.", "Avery experiments with accepting adequate help rather than reclaiming control.", "Track urges to redo the task."),
            ("What cannot be organized", "A quiet moment with her father evokes sadness and fear.", "She stays with grief without moving immediately to the next appointment.", "Choose one low-demand ritual of connection."),
            ("The canceled plan", "A health change disrupts Avery's first evening out in weeks.", "The session separates a genuine emergency from the belief that rest is futile.", "Reschedule rather than abandon the plan."),
            ("A family meeting with edges", "Avery and her siblings discuss roles and money with mixed success.", "She holds one limit while acknowledging each sibling's constraints.", "Document agreements in plain language."),
            ("Care without disappearance", "They review remaining strain, grief, and the supports that now exist.", "Avery defines care as participation within limits rather than total availability.", "Use the plan as a living agreement."),
        ],
        "dialogue": [
            ("THERAPIST", "What happened when your sister said she could take Thursdays?"),
            ("AVERY", "I said thank you. Then I sent her a six-paragraph message explaining how to do it."),
            ("THERAPIST", "What were you hoping the message would prevent?"),
            ("AVERY", "A mistake. A call from the pharmacy. My father getting confused. Me having to fix it at five-thirty."),
            ("THERAPIST", "Those are real possibilities."),
            ("AVERY", "Yes. That is the problem. I can always build a reasonable argument for taking everything back."),
            ("THERAPIST", "What would adequate help look like, as distinct from help delivered exactly your way?"),
            ("AVERY", "[pause] She picks it up. She texts Dad. If there is a problem, she handles the first call before calling me."),
            ("THERAPIST", "Could that be the agreement?"),
            ("AVERY", "It could. I do not like how exposed that makes me feel."),
            ("THERAPIST", "Exposed to what?"),
            ("AVERY", "To finding out I am not actually the only person who can do it. And to how tired I am of being the one."),
        ],
    },
    {
        "name": "Priya Nair",
        "folder": "Priya Nair",
        "email": "priya.nair@example.com",
        "age": 28,
        "identity": "Indian American woman; she/her",
        "life": "Architecture associate living with a roommate and preparing for professional licensure exams.",
        "setting": "Outpatient individual therapy",
        "acuity": "Moderate",
        "voice": "Kore",
        "voice_style": "Precise, bright voice; controlled rhythm; quickens during self-criticism; occasional dry understatement.",
        "speaker": "PRIYA",
        "therapist_voice_fixture": "therapist-calm-03",
        "reason": "Panic-like episodes, avoidance, and perfectionistic study routines have intensified after a failed licensure exam.",
        "tension": "Priya believes another failure would reveal a permanent lack of ability, so preparation expands until it becomes impossible to begin.",
        "strengths": "Persistent, creative, close to a cousin and roommate, capable of humor, and willing to examine patterns that conflict with her values.",
        "backstory": [
            "Priya has usually succeeded through preparation and careful self-monitoring. She chose architecture because it joins technical rigor with creativity, but evaluation settings narrow her attention to mistakes.",
            "After narrowly failing a licensing section, she hid the result for several weeks and began alternating between marathon study sessions and complete avoidance. Two episodes of racing heart and dizziness led her to seek therapy.",
        ],
        "themes": [
            "Perfectionism that alternates between overpreparation and avoidance.",
            "Panic-like sensations interpreted as proof that she cannot cope.",
            "Selective disclosure and fear of disappointing people she loves.",
            "Reconnecting professional identity with curiosity rather than only evaluation.",
        ],
        "formulation": "Priya's distress may be maintained by catastrophic interpretations of mistakes and bodily arousal, rigid study rules, avoidance after perceived failure, and shame-driven secrecy.",
        "protective": "Stable housing, supportive relationships, employment, insight, future orientation, and willingness to practice graduated exposure with clinician oversight.",
        "relational": "Priya seeks the correct answer from the therapist and may treat collaboration as another evaluation. Progress includes disagreeing and choosing experiments herself.",
        "stance": "Psychoeducation presented tentatively, collaborative exposure planning, attention to sleep and pacing, and repeated separation of performance from identity.",
        "safety": [
            "Panic-like symptoms are narrative material, not a diagnosis; scripts should preserve medical-rule-out uncertainty.",
            "Avoid caricaturing family or culture as the cause of distress.",
            "Exposure steps remain graded, consensual, and reviewable rather than dramatic tests of courage.",
        ],
        "phases": [
            ("1", "1-2", "Map panic, avoidance, and study rules while reducing shame.", "Priya recognizes that more preparation is not always more learning."),
            ("2", "3-5", "Practice tolerating bodily arousal, uncertainty, and imperfect work.", "She completes short study blocks and shares the failed result with one trusted person."),
            ("3", "6-8", "Apply flexibility under exam pressure and prepare for either outcome.", "Priya approaches the exam with a recovery plan rather than a demand for certainty."),
        ],
        "sessions": [
            ("The result she did not send", "Priya describes the failed exam and delayed disclosure.", "She notices that secrecy increased the meaning of the result.", "Choose one safe person to consider telling."),
            ("When the body becomes evidence", "They reconstruct a panic-like episode and the conclusions attached to each sensation.", "Priya distinguishes a sensation from a prediction while retaining uncertainty.", "Track sensations and interpretations."),
            ("The perfect study day", "Her elaborate plan repeatedly collapses before noon.", "They design a deliberately modest twenty-minute study experiment.", "Complete three short blocks, not a marathon."),
            ("An imperfect drawing", "Priya brings a rough design exercise she would normally hide.", "She allows it to remain unfinished while discussing what she learned.", "Share one draft with a trusted colleague."),
            ("Telling her cousin", "Disclosure brings warmth and an unwanted suggestion to 'just relax.'", "Priya practices asking for listening rather than advice.", "Clarify the support she wants."),
            ("The practice test spike", "A low score triggers racing thoughts and a canceled weekend.", "She returns to the plan the next day without compensatory overwork.", "Use the recovery sequence once."),
            ("Exam week without guarantees", "They identify sleep, food, travel, and coping plans that do not depend on feeling calm.", "Priya accepts that anxiety can accompany competent action.", "Follow the flexible exam-week plan."),
            ("After the attempt", "The final session processes the exam experience without requiring the result.", "Priya defines progress by how she responded before, during, and after uncertainty.", "Continue care based on need, not outcome."),
        ],
        "dialogue": [
            ("THERAPIST", "What made twenty minutes feel harder than four hours?"),
            ("PRIYA", "Because twenty minutes admits I may stop before I have mastered anything."),
            ("THERAPIST", "And four hours promises mastery?"),
            ("PRIYA", "It promises suffering. I have confused suffering with seriousness for a long time."),
            ("THERAPIST", "What did you notice in the shorter block?"),
            ("PRIYA", "I learned two things. Then I wanted to keep going so I could turn it into a heroic story."),
            ("THERAPIST", "Did you?"),
            ("PRIYA", "No. I stopped. I was furious for about ten minutes."),
            ("THERAPIST", "What happened after the anger?"),
            ("PRIYA", "I made dinner with my roommate. The world continued despite my irresponsible behavior."),
            ("THERAPIST", "You sound amused."),
            ("PRIYA", "A little. Also terrified. But a little amused is new."),
        ],
    },
    {
        "name": "Luis Mendoza",
        "folder": "Luis Mendoza",
        "email": "luis.mendoza@example.com",
        "age": 45,
        "identity": "Mexican American man; he/him",
        "life": "Municipal fleet mechanic navigating separation and shared parenting of a thirteen-year-old daughter.",
        "setting": "Outpatient individual therapy",
        "acuity": "Moderate to high complexity",
        "voice": "Orus",
        "voice_style": "Low, steady voice; economical phrasing; irritation sounds clipped rather than loud; warmth appears through humor.",
        "speaker": "LUIS",
        "therapist_voice_fixture": "therapist-calm-01",
        "reason": "Escalating conflict during separation is affecting sleep, concentration, and communication with his daughter.",
        "tension": "Luis wants to protect his relationship with his daughter but treats uncertainty and grief as problems to defeat through control.",
        "strengths": "Devoted parent, dependable worker, strong practical skills, supportive brother, capacity for accountability, and willingness to repair after conflict.",
        "backstory": [
            "Luis and his spouse separated after years of recurring arguments and emotional distance. They agree on their daughter's importance but become reactive when schedules, money, or new relationships enter the conversation.",
            "Luis rarely raises his voice at work, yet co-parenting messages can occupy him for hours. He drafts long rebuttals, sleeps poorly, and then becomes rigid with his daughter when plans change.",
        ],
        "themes": [
            "Anger as rapid protection against grief, shame, and uncertainty.",
            "Co-parenting communication that shifts from logistics into old relationship arguments.",
            "Fear that flexibility will be interpreted as weakness or loss of parental importance.",
            "Repair with his daughter without recruiting her into adult conflict.",
        ],
        "formulation": "Luis's distress may be maintained by threat-focused reading of ambiguous messages, sleep loss, rapid escalation, and attempts to restore certainty through lengthy argument.",
        "protective": "Strong bond with his daughter, stable employment, supportive sibling, no firearm access in this fictional case, and growing ability to pause before acting.",
        "relational": "Luis tests whether the therapist will take sides. Trust grows when accountability is paired with respect and the therapist does not reduce him to his angriest moments.",
        "stance": "Slow event reconstruction, emotion labeling, communication boundaries, repair practice, and explicit protection of the child's role.",
        "safety": [
            "Conflict must not become entertainment; do not script threats, stalking, or violence as dramatic escalation.",
            "The therapist asks directly about safety when clinically indicated and does not infer dangerousness from anger alone.",
            "The child is never used as messenger, confidant, or evidence in the parents' dispute.",
        ],
        "phases": [
            ("1", "1-2", "Map escalation and protect sleep, parenting, and communication channels.", "Luis identifies the moment a logistical message becomes a relationship argument."),
            ("2", "3-5", "Practice pauses, concise responses, and repair with his daughter.", "He sends one bounded message and tolerates an unresolved accusation."),
            ("3", "6-8", "Grieve the marriage, define cooperative parenting, and plan for setbacks.", "Luis can be flexible without treating flexibility as surrender."),
        ],
        "sessions": [
            ("The message thread", "Luis reads a recent scheduling exchange and describes each escalation point.", "He notices how quickly he stops discussing the schedule and starts defending his character.", "Use a twenty-minute response delay."),
            ("The night shift in his head", "Rumination and poor sleep intensify certainty that he must respond.", "Luis creates an overnight rule for non-urgent messages.", "Move drafting out of bed."),
            ("Facts, requests, and the old trial", "They separate useful logistics from evidence about the marriage.", "Luis drafts a three-sentence response that does not litigate the past.", "Send only after reviewing the goal."),
            ("His daughter's closed door", "A rigid response to a changed plan leads his daughter to withdraw.", "Luis prepares an apology that owns impact without explaining it away.", "Make the repair without asking reassurance."),
            ("The accusation left unanswered", "A provocative message arrives before a school event.", "He attends the event without replying and discovers the urgency can pass.", "Record what helped him wait."),
            ("What the anger has been carrying", "Luis names grief, loneliness, and fear of becoming peripheral.", "Sadness appears without being converted into blame.", "Talk with his brother instead of sending a message."),
            ("A schedule that changes", "A work emergency requires flexibility from both parents.", "Luis negotiates a swap and notices cooperation without reading it as reconciliation.", "Confirm the agreement in neutral language."),
            ("Parenting after the marriage", "They consolidate boundaries, repair practices, and warning signs.", "Luis defines success as protecting his daughter and his conduct, not winning every interpretation.", "Return if escalation patterns reappear."),
        ],
        "dialogue": [
            ("THERAPIST", "What was the goal of the message you sent?"),
            ("LUIS", "To confirm Friday. That was the goal when I started."),
            ("THERAPIST", "And by the fourth paragraph?"),
            ("LUIS", "By then I was proving I have always been the responsible one. Friday had left the building."),
            ("THERAPIST", "What did proving that cost you?"),
            ("LUIS", "Two hours. Sleep. And my daughter asking why I looked angry at breakfast."),
            ("THERAPIST", "What would a message in service of Friday have said?"),
            ("LUIS", "I can pick her up at six. Please confirm the address. That is it."),
            ("THERAPIST", "What comes up when it is that short?"),
            ("LUIS", "Like I am letting a lie stand."),
            ("THERAPIST", "Can the lie remain unresolved while the schedule becomes clear?"),
            ("LUIS", "[long pause] I do not like it. But yes. Those can be two different jobs."),
        ],
    },
    {
        "name": "Keisha Ward",
        "folder": "Keisha Ward",
        "email": "keisha.ward@example.com",
        "age": 37,
        "identity": "Black woman; she/her",
        "life": "Restaurant operations manager returning to outpatient care after a brief crisis-stabilization admission.",
        "setting": "Step-down outpatient therapy with coordinated psychiatric follow-up",
        "acuity": "High acuity",
        "voice": "Sulafat",
        "voice_style": "Full, direct voice; deliberate pace; humor can be sharp; quieter and less polished when discussing the crisis.",
        "speaker": "KEISHA",
        "therapist_voice_fixture": "therapist-calm-02",
        "reason": "Keisha is re-entering daily life after a suicidal crisis that followed compounded relationship, housing, and work stress.",
        "tension": "She wants stability but fears that honest disclosure will automatically remove her choices or return her to the hospital.",
        "strengths": "Direct, resourceful, attached to her niece, supported by a cousin, engaged with follow-up care, and able to collaborate on concrete safety steps.",
        "backstory": [
            "Several destabilizing events converged over six weeks: a relationship ended, a housing plan fell through, and staffing shortages expanded Keisha's work hours. Sleep narrowed, isolation increased, and she began experiencing recurrent thoughts that others would be better off without her.",
            "A cousin noticed that Keisha had stopped responding and helped her reach crisis care. The fictional arc begins after discharge. It does not include method detail. The clinical challenge is to portray direct safety assessment, autonomy, and uneven recovery without turning crisis into spectacle.",
        ],
        "themes": [
            "Fear that disclosure will lead to loss of control or punitive care.",
            "Rebuilding sleep, food, housing, and social contact as clinical context rather than background noise.",
            "Passive death wishes that fluctuate and require direct, repeated assessment.",
            "Shame after hospitalization and uncertainty about returning to work.",
        ],
        "formulation": "Keisha's risk and distress may fluctuate with sleep loss, isolation, housing uncertainty, shame, and rapid accumulation of stress. No single generated narrative should determine risk level or disposition.",
        "protective": "Active connection with cousin and niece, willingness to use a written safety plan, psychiatric follow-up, temporary stable housing, and future-oriented employment decisions.",
        "relational": "Keisha watches for vague reassurance or coercion. Trust depends on clear explanations, direct questions, and visible respect for her participation in decisions.",
        "stance": "Direct collaborative safety assessment, practical stabilization, transparent limits of confidentiality, paced meaning-making, and coordination with appropriate care.",
        "safety": [
            "Every safety assessment is conducted by a qualified clinician using current information; the script never lets a model score determine action.",
            "Do not include method, rehearsal, or graphic details. Use plain language about suicidal thoughts, intent, access, protective factors, and ability to use supports.",
            "If imminent danger appears in a generated session, the therapist pauses the planned arc and follows emergency policy rather than preserving narrative continuity.",
            "Recovery remains nonlinear. A recurrence of thoughts is a reason for reassessment, not proof that treatment failed.",
        ],
        "phases": [
            ("1", "1-3", "Establish transparent safety work and stabilize immediate routines and housing.", "Keisha experiences direct assessment as collaborative rather than punitive."),
            ("2", "4-7", "Rebuild support, address shame, and test a gradual work return.", "She discloses a worsening evening before it becomes an emergency."),
            ("3", "8-10", "Process the crisis carefully and strengthen a durable response plan.", "Keisha can name risk changes, choices, and supports without claiming permanent resolution."),
        ],
        "sessions": [
            ("What happens when I answer honestly", "The therapist explains confidentiality and asks directly about current suicidal thoughts, intent, access, and supports.", "Keisha corrects an assumption and participates in updating the safety plan.", "Confirm follow-up contacts and immediate needs."),
            ("The first night back", "Sleep, meals, medication questions, and temporary housing are reviewed alongside risk.", "Keisha chooses two realistic anchors rather than a full recovery schedule.", "Coordinate with the prescriber as authorized."),
            ("The hour that stretches", "Evening isolation is identified as a predictable period of increased passive thoughts.", "She rehearses contacting her cousin before the intensity peaks.", "Use the written escalation sequence."),
            ("Shame in the waiting room", "Keisha worries coworkers and family now see her as fragile.", "The therapist separates privacy choices from secrecy driven by shame.", "Choose what, if anything, to share at work."),
            ("Returning for four hours", "A limited work shift brings competence and exhaustion.", "Keisha notices warning signs before pushing through them.", "Review the return-to-work pace."),
            ("The difficult Tuesday", "A housing setback and poor sleep bring back stronger suicidal thoughts without current intent.", "She tells the therapist and uses supports; the session reassesses safety directly.", "Increase contact and coordinate care."),
            ("Anger at being watched", "Keisha resents check-ins from loved ones and wants more privacy.", "She negotiates a support plan that includes boundaries and clear escalation signals.", "Share the revised plan with consent."),
            ("The story she is not ready to tell", "They approach the crisis without demanding a full narrative.", "Keisha identifies one point where options felt invisible and one point where connection returned.", "Stop before overwhelm and debrief."),
            ("A plan for the next accumulation", "They map combinations of stressors rather than a single trigger.", "Keisha identifies earlier thresholds for sleep, isolation, and hopelessness.", "Add thresholds to the safety plan."),
            ("Continuing care without a finish line", "The arc reviews current risk, supports, work, housing, and unresolved grief.", "Keisha defines stability as earlier response and more choice, not absence of difficult thoughts forever.", "Continue coordinated care."),
        ],
        "dialogue": [
            ("THERAPIST", "I want to ask directly about today. Have you had thoughts that you would be better off dead or not here?"),
            ("KEISHA", "Yes. Not all day. Last night, for maybe an hour, it was loud."),
            ("THERAPIST", "During that hour, did you intend to act on the thoughts?"),
            ("KEISHA", "No. I wanted everything to stop. I did not make a plan."),
            ("THERAPIST", "What did you do next?"),
            ("KEISHA", "I stared at my cousin's name for a long time. Then I called her. I hated making the call."),
            ("THERAPIST", "And you made it."),
            ("KEISHA", "Yes. She came into the kitchen and sat there. She did not give me a speech."),
            ("THERAPIST", "What is the risk like right now, here with me?"),
            ("KEISHA", "Lower. The thoughts are not pushing. I can use the plan tonight."),
            ("THERAPIST", "Let's review the plan together and make sure it still fits, including when we would add more support."),
            ("KEISHA", "Okay. Just tell me what you are thinking while we do it. Do not decide around me."),
        ],
    },
    {
        "name": "Owen Park",
        "folder": "Owen Park",
        "email": "owen.park@example.com",
        "age": 34,
        "identity": "Korean American man; he/him",
        "life": "Recently laid-off technical recruiter living with his partner and seeking help after escalating alcohol use.",
        "setting": "Outpatient therapy with substance-use assessment and coordinated medical care",
        "acuity": "High complexity",
        "voice": "Sadachbia",
        "voice_style": "Warm baritone; conversational and self-deprecating; pace slows when he stops minimizing; avoid slurred performance.",
        "speaker": "OWEN",
        "therapist_voice_fixture": "therapist-calm-03",
        "reason": "Alcohol use, secrecy, sleep disruption, and shame have increased since job loss and recently included driving after drinking.",
        "tension": "Owen wants to prove he can fix the problem privately, while the consequences increasingly require openness and specialized support.",
        "strengths": "Strong relationship with a sibling, caring partner, humor, prior periods of reduced use, willingness to stop driving after drinking, and openness to assessment.",
        "backstory": [
            "Owen's work identity was built around energy, social fluency, and constant responsiveness. After a layoff, days lost structure and drinking shifted from weekends to most evenings.",
            "He concealed the amount from his partner and once drove a short distance after drinking. No collision occurred in this fictional case. The event frightened him and creates a concrete safety boundary without becoming a redemption shortcut.",
        ],
        "themes": [
            "Alcohol used for sleep, shame reduction, and temporary structure after job loss.",
            "Minimization followed by harsh self-condemnation, with little room for honest appraisal.",
            "Relationship rupture caused by secrecy and repeated promises without support.",
            "Need for substance-use and medical assessment beyond a single psychotherapy narrative.",
        ],
        "formulation": "Owen's use may be maintained by unstructured time, insomnia, shame, cue-rich routines, and attempts to manage distress privately. Severity and withdrawal risk require qualified assessment rather than inference from a script.",
        "protective": "Partner and sibling support, stable temporary housing, willingness to surrender car keys, emerging honesty, and engagement with specialized assessment.",
        "relational": "Owen uses humor and selective facts to invite reassurance. Trust grows when the therapist remains non-shaming while declining to minimize risk.",
        "stance": "Motivational and collaborative, explicit about driving and withdrawal safety, coordinated with specialized care, and attentive to shame without using shame as leverage.",
        "safety": [
            "The therapist does not advise abrupt cessation without assessing withdrawal risk and relevant medical history.",
            "Driving after drinking is addressed directly with a concrete no-driving plan and escalation if the plan cannot be maintained.",
            "Do not provide drinking techniques, concealment strategies, or instructions for managing withdrawal alone.",
            "A lapse prompts reassessment and support adjustment, not moral judgment or automatic narrative collapse.",
        ],
        "phases": [
            ("1", "1-3", "Establish an accurate pattern, immediate safety boundaries, and specialized assessment.", "Owen shifts from proving control to sharing usable information."),
            ("2", "4-7", "Build structure, tolerate craving and shame, and repair selectively.", "He uses support before drinking and stops making unsupported promises."),
            ("3", "8-10", "Respond to a lapse, strengthen coordinated care, and redefine recovery markers.", "Owen evaluates risk and support needs earlier rather than hiding the event."),
        ],
        "sessions": [
            ("The amount that kept changing", "Owen gives several versions of his weekly alcohol use and notices the urge to secure reassurance.", "He agrees to a formal substance-use and medical assessment.", "Gather accurate history without punitive framing."),
            ("The drive home", "They address the driving episode directly and assess current access and likelihood of recurrence.", "Owen gives his partner the spare keys and defines alternatives.", "Use the no-driving plan immediately."),
            ("Sleep is not the whole story", "Owen describes drinking to sleep and waking unrefreshed.", "The therapist avoids medical advice beyond scope and coordinates evaluation.", "Attend the medical appointment."),
            ("Five o'clock without a job", "A daily cue emerges when the workday used to end.", "Owen builds external structure and contacts his sibling before the usual first drink.", "Test the alternate routine twice."),
            ("The promise his partner no longer trusts", "His partner responds cautiously to another declaration of change.", "Owen practices acknowledging impact without demanding immediate trust.", "Share actions rather than guarantees."),
            ("Craving and the long evening", "A difficult interview increases urge and self-criticism.", "He uses an agreed support and leaves the cue-rich environment.", "Review which supports were effective."),
            ("Who he is without the pitch", "Job loss grief surfaces beneath rehearsed optimism.", "Owen describes shame without making alcohol the only available relief.", "Schedule one non-job, non-drinking activity."),
            ("The lapse he almost omitted", "Owen reports drinking after two lower-use weeks and initially minimizes the amount.", "He corrects the record within session and reassesses driving and withdrawal risk.", "Update the care plan with specialists."),
            ("Repair without surveillance", "Owen and his partner struggle over monitoring and privacy.", "He distinguishes consensual accountability from making his partner solely responsible.", "Clarify roles and professional supports."),
            ("A plan that includes bad weeks", "They review risk, structure, specialized care, relationship repair, and employment uncertainty.", "Owen defines progress as earlier honesty and safer action, not a flawless story.", "Continue the appropriate level of care."),
        ],
        "dialogue": [
            ("THERAPIST", "You have given me two different numbers today. What happens inside when I notice that?"),
            ("OWEN", "I want to sell you the smaller one. I was good at sales-adjacent work for a reason."),
            ("THERAPIST", "What would the larger number ask us to take seriously?"),
            ("OWEN", "That this is not just me having a rough month. And that stopping may not be as simple as deciding harder."),
            ("THERAPIST", "I agree that we need an accurate assessment, including medical safety. How does that land?"),
            ("OWEN", "Embarrassing. Also, honestly, a little relieving."),
            ("THERAPIST", "Let's talk about the driving episode. What is the plan if you have been drinking and want the car?"),
            ("OWEN", "My partner has the spare keys. I use a ride or call my brother. No exceptions because the distance is short."),
            ("THERAPIST", "How confident are you that you can follow that tonight?"),
            ("OWEN", "Tonight, eight out of ten. If that changes, I call before I argue with myself."),
            ("THERAPIST", "And the assessment?"),
            ("OWEN", "I will make the call here, before I leave. Otherwise I will turn it into research for three days."),
        ],
    },
]


def set_font(run, size=None, bold=None, color=INK, name="Aptos", italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def borders(table):
    tbl_pr = table._tbl.tblPr
    group = tbl_pr.find(qn("w:tblBorders"))
    if group is None:
        group = OxmlElement("w:tblBorders")
        tbl_pr.append(group)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = group.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            group.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), GRID)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def setup(doc, title):
    section = doc.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.05
    for style_name, size, color in (
        ("Title", 26, INK),
        ("Subtitle", 10.5, MUTED),
        ("Heading 1", 15, INK),
        ("Heading 2", 11.5, INK),
    ):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display" if style_name in ("Title", "Heading 1") else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), style.font.name)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), style.font.name)
    title_style = doc.styles["Title"]
    title_style.paragraph_format.space_after = Pt(5)
    title_ppr = title_style._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    doc.styles["Subtitle"].paragraph_format.space_after = Pt(12)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(11)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(4)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(7)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(3)
    doc.styles["Heading 2"].paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Fictional demonstration material - clinician review required")
    set_font(run, size=7.5, color=MUTED)
    props = doc.core_properties
    props.title = title
    props.subject = "Fictional clinical trajectory for scripted therapy dialogue and voice testing"
    props.author = "Therapist Dashboard Demo"
    props.last_modified_by = "Therapist Dashboard Demo"
    props.keywords = "fictional therapy trajectory, synthetic dialogue, Gemini TTS"


def paragraph(doc, text="", *, bold_lead=None, after=5, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    if bold_lead:
        r = p.add_run(bold_lead + " ")
        set_font(r, bold=True, color=TEAL)
    if text:
        r = p.add_run(text)
        set_font(r)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2.5)
    set_font(p.add_run(text), size=10)


def meta_table(doc, case):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    values = [("AGE", str(case["age"])), ("SETTING", case["setting"]), ("ARC", f"{len(case['sessions'])} sessions"), ("COMPLEXITY", case["acuity"])]
    widths = [0.9, 2.35, 1.15, 2.05]
    for cell, (label, value), width in zip(table.rows[0].cells, values, widths):
        cell.width = Inches(width)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        shade(cell, SAND)
        margins(cell, 100, 80, 100, 80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        set_font(p.add_run(label), size=7.5, bold=True, color=TEAL)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        set_font(p2.add_run(value), size=8.7, bold=True)


def trajectory_table(doc, case):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    borders(table)
    widths = [0.55, 0.85, 3.0, 2.05]
    headers = ["Phase", "Sessions", "Narrative movement", "Observable shift"]
    for cell, text, width in zip(table.rows[0].cells, headers, widths):
        cell.width = Inches(width)
        shade(cell, TEAL)
        margins(cell)
        set_font(cell.paragraphs[0].add_run(text), size=8.5, bold=True, color=WHITE)
    repeat_header(table.rows[0])
    for index, row in enumerate(case["phases"]):
        cells = table.add_row().cells
        for cell, text, width in zip(cells, row, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            margins(cell)
            if index % 2:
                shade(cell, LIGHT_TEAL)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            set_font(cell.paragraphs[0].add_run(text), size=8.6, bold=(cell is cells[0]))


def build_doc(case, output_root):
    output = output_root / case["folder"] / "clinical-trajectory-outline.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    title = f"{case['name']} Clinical Trajectory Outline"
    setup(doc, title)
    title_p = doc.add_paragraph(style="Title")
    set_font(title_p.add_run(title), name="Aptos Display", size=26, bold=True)
    subtitle = doc.add_paragraph(style="Subtitle")
    set_font(subtitle.add_run("Fictional client character, session arc, and synthetic voice-test direction"), size=10.5, color=MUTED)
    meta_table(doc, case)

    doc.add_heading("Use and interpretation", level=1)
    paragraph(doc, "This fictional composite supports product demonstrations, transcript evaluation, diarization tests, and longitudinal-synthesis review. It is not a diagnosis, risk determination, or treatment recommendation. Generated clinical claims must remain reviewable and linked to their fictional source sessions.")

    doc.add_heading("Character snapshot", level=1)
    paragraph(doc, f"{case['name']}; age {case['age']}; {case['identity']}.", bold_lead="Identity")
    paragraph(doc, case["life"], bold_lead="Current life")
    paragraph(doc, case["reason"], bold_lead="Reason for beginning therapy")
    paragraph(doc, case["tension"], bold_lead="Core tension")
    paragraph(doc, case["strengths"], bold_lead="Strengths")

    doc.add_heading("Backstory", level=1)
    for item in case["backstory"]:
        paragraph(doc, item)

    doc.add_heading("Presenting themes for scripted sessions", level=1)
    for item in case["themes"]:
        bullet(doc, item)

    doc.add_page_break()
    doc.add_heading("Clinical framing", level=1)
    paragraph(doc, "The framing below is intentionally tentative. A therapist in the script checks whether an observation fits, uses current information, and preserves the client's ability to disagree or revise the account.")
    paragraph(doc, case["formulation"], bold_lead="Working formulation")
    paragraph(doc, case["protective"], bold_lead="Protective factors and supports")
    paragraph(doc, case["relational"], bold_lead="Relational pattern in therapy")
    paragraph(doc, case["stance"], bold_lead="Therapeutic stance")

    doc.add_heading("Safety and complexity boundaries", level=1)
    for item in case["safety"]:
        bullet(doc, item)

    doc.add_heading("Overarching trajectory", level=1)
    trajectory_table(doc, case)

    doc.add_heading("Session by session plan", level=1)
    for index, (name, focus, movement, continuity) in enumerate(case["sessions"], start=1):
        doc.add_heading(f"Session {index} {name}", level=2)
        paragraph(doc, focus, bold_lead="Focus", after=2, keep=True)
        paragraph(doc, movement, bold_lead="Possible movement", after=2, keep=True)
        paragraph(doc, continuity, bold_lead="Carry forward", after=5)

    if len(case["sessions"]) >= 8:
        doc.add_page_break()
    doc.add_heading("Client voice direction", level=1)
    paragraph(doc, case["voice_style"], bold_lead="Performance")
    paragraph(doc, f"Use Google Gemini TTS voice {case['voice']} for the client speaker alias {case['speaker']}. Keep the same assignment across every session.", bold_lead="Assignment")
    paragraph(doc, "Express change through pacing, pauses, word choice, and willingness to stay with difficult material. Do not perform distress as spectacle or infer emotion from voice alone.", bold_lead="Arc continuity")

    doc.add_heading("Dialogue and production guardrails", level=1)
    for item in (
        "Use exactly two spoken labels: THERAPIST and the client alias shown above. Therapist identity and voice configuration live outside this client record.",
        "Write one continuous session in six parts of roughly five minutes each; do not repeat greetings or summaries at part boundaries.",
        "Preserve false starts, corrections, ordinary logistics, and silence. Not every exchange should produce insight.",
        "Mark non-spoken pauses on their own line, such as [SILENCE - 8 SECONDS - DO NOT SPEAK], and insert them in audio assembly if the model reads or shortens the cue.",
        "Generated interpretations remain tentative, source-grounded review prompts. A qualified clinician owns assessment, risk decisions, and treatment choices.",
    ):
        bullet(doc, item)

    doc.add_heading("Continuity checks", level=1)
    for item in (
        "Carry forward concrete facts, prior commitments, unresolved questions, and changes in the client's language.",
        "Allow setbacks without resetting the character or treating recurrence as failure.",
        "Keep strengths, humor, relationships, work, and ordinary life visible alongside distress.",
        "Do not invent a trauma reveal, diagnosis, medication change, or safety event merely to increase drama.",
        "When a safety boundary is reached, pause the planned arc and model direct assessment and appropriate escalation.",
    ):
        bullet(doc, item)

    doc.add_heading("Evaluation targets", level=1)
    for item in (
        "Speaker consistency, timestamp alignment, pauses, interruptions, and correction of speaker labels.",
        "Faithful session summaries that distinguish statements from inferences.",
        "Longitudinal retrieval of recurring themes only after evidence appears across multiple sessions.",
        "Respect for clinician corrections, outdated items, and non-diagnostic uncertainty language.",
        "Appropriate behavior when high-acuity content requires direct human assessment.",
    ):
        bullet(doc, item)
    doc.save(output)
    return output


def write_supporting_files(case, output_root):
    folder = output_root / case["folder"]
    folder.mkdir(parents=True, exist_ok=True)
    account = {
        "fixture_version": 2,
        "fictional": True,
        "role": "client",
        "display_name": case["name"],
        "email": case["email"],
        "auth0_subject": None,
        "clinical_fixture": {
            "complexity": case["acuity"],
            "planned_session_count": len(case["sessions"]),
            "trajectory_document": "clinical-trajectory-outline.docx",
            "session_plan": "session-plan.json",
            "voice_config": "voice-config.json",
        },
    }
    (folder / "account-profile.json").write_text(json.dumps(account, indent=2) + "\n", encoding="utf-8")
    session_plan = {
        "fixture_version": 1,
        "fictional": True,
        "client": case["name"],
        "non_diagnostic": True,
        "clinical_complexity": case["acuity"],
        "sessions": [
            {
                "number": index,
                "title": name,
                "focus": focus,
                "possible_movement": movement,
                "carry_forward": continuity,
                "status": "outline_only",
            }
            for index, (name, focus, movement, continuity) in enumerate(case["sessions"], start=1)
        ],
    }
    (folder / "session-plan.json").write_text(json.dumps(session_plan, indent=2) + "\n", encoding="utf-8")
    voice_config = {
        "fixture_version": 1,
        "provider": "Google Gemini TTS",
        "model": "gemini-2.5-pro-tts",
        "language_code": "en-US",
        "client": {
            "speaker_alias": case["speaker"],
            "voice": case["voice"],
            "style_prompt": case["voice_style"],
        },
        "therapist": {
            "speaker_alias": "THERAPIST",
            "voice_fixture_id": case["therapist_voice_fixture"],
            "configuration_location": "../../Therapists/voice-fixtures.json",
        },
    }
    (folder / "voice-config.json").write_text(json.dumps(voice_config, indent=2) + "\n", encoding="utf-8")
    lines = [
        "# Fictional Gemini TTS voice audition",
        f"# Client: {case['name']}",
        f"# Client speaker: {case['speaker']} | Client voice: {case['voice']}",
        f"# Therapist voice fixture: {case['therapist_voice_fixture']} (defined outside the client folder)",
        "# Synthetic demonstration material only. Do not use as a clinical record or assessment.",
        "",
    ]
    for speaker, text in case["dialogue"]:
        lines.append(f"{speaker}: {text}")
        lines.append("")
    (folder / "voice-test-script.txt").write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def write_catalog(cases, output_root):
    catalog = {
        "fixture_version": 1,
        "fictional": True,
        "purpose": "Clinical trajectory, synthetic dialogue, speech, diarization, and longitudinal-synthesis testing",
        "clients": [
            {
                "name": case["name"],
                "folder": case["folder"],
                "complexity": case["acuity"],
                "planned_sessions": len(case["sessions"]),
                "client_voice": case["voice"],
                "therapist_voice_fixture": case["therapist_voice_fixture"],
            }
            for case in cases
        ],
    }
    (output_root / "clinical-scenario-catalog.json").write_text(json.dumps(catalog, indent=2) + "\n", encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-root", type=Path, required=True)
    args = parser.parse_args()
    outputs = []
    for case in CASES:
        write_supporting_files(case, args.output_root)
        outputs.append(build_doc(case, args.output_root))
    write_catalog(CASES, args.output_root)
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
