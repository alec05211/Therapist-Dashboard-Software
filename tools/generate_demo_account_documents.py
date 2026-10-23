"""Generate fictional client and therapist account documents for demo-data."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLACK = "000000"
NAVY = "28465E"
PALE_BLUE = "EEF4F7"
LIGHT_GRAY = "D9D9D9"
TEXT_GRAY = "4A5560"


THERAPISTS = [
    {
        "folder": "Nadia Brooks",
        "name": "Nadia Brooks",
        "credentials": "LCSW",
        "email": "nadia.brooks@example.com",
        "practice": "Harborline Counseling Group",
        "jurisdiction": "New York",
        "license_id": "DEMO-NY-LCSW-10427",
        "specialties": "Adult anxiety, workplace stress, and life transitions",
        "modality": "Individual telehealth and in-person sessions",
        "bio": (
            "Nadia uses a collaborative, practical approach that helps adults notice patterns, "
            "build steadier routines, and make room for values-based choices."
        ),
        "client": "Jordan Kim",
    },
    {
        "folder": "Samuel Ortiz",
        "name": "Samuel Ortiz",
        "credentials": "LMFT",
        "email": "samuel.ortiz@example.com",
        "practice": "North Grove Therapy",
        "jurisdiction": "Pennsylvania",
        "license_id": "DEMO-PA-LMFT-20816",
        "specialties": "Relationships, caregiver strain, and adjustment",
        "modality": "Individual, couples, and telehealth sessions",
        "bio": (
            "Samuel helps clients slow difficult interaction cycles, name competing needs, "
            "and practice communication that is both direct and compassionate."
        ),
        "client": "Avery Chen",
    },
]


CLIENTS = [
    {
        "folder": "Jordan Kim",
        "name": "Jordan Kim",
        "email": "jordan.kim@example.com",
        "reason": "Work pressure has begun to affect sleep, recovery time, and decision making.",
        "goals": [
            "Create a more consistent evening wind-down routine",
            "Practice clearer boundaries around after-hours work",
            "Notice early signs of overload before reaching exhaustion",
        ],
        "preferences": [
            ("Session format", "Telehealth preferred"),
            ("Appointment time", "Weekday evenings"),
            ("Reminders", "Email 24 hours before a session"),
            ("Communication", "Email for scheduling only; discuss care in session"),
        ],
    },
    {
        "folder": "Avery Chen",
        "name": "Avery Chen",
        "email": "avery.chen@example.com",
        "reason": "Caregiving demands and family communication have made it difficult to maintain personal routines.",
        "goals": [
            "Clarify which caregiving responsibilities can be shared",
            "Prepare for calmer conversations about limits and support",
            "Rebuild one dependable weekly activity for rest",
        ],
        "preferences": [
            ("Session format", "In person preferred; telehealth acceptable"),
            ("Appointment time", "Weekday mornings"),
            ("Reminders", "Email 48 hours before a session"),
            ("Communication", "Email for scheduling only; no clinical details by email"),
        ],
    },
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=75, start=120, bottom=75, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), LIGHT_GRAY)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def configure_document(doc: Document, title: str, subject: str) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.52)
    section.bottom_margin = Inches(0.52)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.02

    title_style = styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style.font.size = Pt(20)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string(BLACK)
    title_style.paragraph_format.space_after = Pt(5)
    title_ppr = title_style._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size, before, after in (
        ("Heading 1", 13.5, 9, 3),
        ("Heading 2", 11.5, 6, 2),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    props = doc.core_properties
    props.title = title
    props.subject = subject
    props.author = "Therapist Dashboard Demo"
    props.last_modified_by = "Therapist Dashboard Demo"
    props.keywords = "fictional demo account"

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Fictional demonstration record - not valid for identity or license verification")
    run.font.name = "Aptos"
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor.from_string(TEXT_GRAY)


def add_title_block(doc: Document, title: str, subtitle: str) -> None:
    paragraph = doc.add_paragraph(style="Title")
    paragraph.add_run(title)
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(7)
    run = intro.add_run(subtitle)
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(TEXT_GRAY)


def add_key_value_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(2.0)
    table.columns[1].width = Inches(4.7)
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    header.cells[0].text = "Field"
    header.cells[1].text = "Demo value"
    for cell in header.cells:
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for index, (label, value) in enumerate(rows):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        if index % 2:
            set_cell_shading(cells[0], PALE_BLUE)
            set_cell_shading(cells[1], PALE_BLUE)
        for cell in cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
        cells[0].paragraphs[0].runs[0].font.bold = True


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run(item)


def add_verification_steps(doc: Document) -> None:
    steps = [
        "Match the legal name and jurisdiction supplied during onboarding.",
        "Confirm the credential through the applicable state licensing board or an approved verification provider.",
        "Record the result, source, verification date, and reviewer in the application database.",
        "Do not grant therapist permissions until the application marks the credential as verified.",
    ]
    for index, step in enumerate(steps, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        paragraph.paragraph_format.space_after = Pt(3)
        lead = paragraph.add_run(f"{index}. ")
        lead.font.bold = True
        paragraph.add_run(step)


def build_therapist_document(output_root: Path, therapist: dict[str, str]) -> Path:
    output = output_root / "Therapists" / therapist["folder"] / "therapist-profile-and-verification.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    title = f"{therapist['name']} Therapist Profile"
    configure_document(doc, title, "Fictional therapist account and credential-verification example")
    add_title_block(
        doc,
        title,
        "A fictional profile for testing therapist onboarding, practice membership, and credential-review workflows.",
    )
    doc.add_heading("Account and practice details", level=1)
    add_key_value_table(
        doc,
        [
            ("Professional name", f"{therapist['name']}, {therapist['credentials']}"),
            ("Role", "Therapist"),
            ("Sign in email", therapist["email"]),
            ("Practice", therapist["practice"]),
            ("Primary jurisdiction", therapist["jurisdiction"]),
            ("Session format", therapist["modality"]),
        ],
    )
    doc.add_heading("Professional profile", level=1)
    doc.add_paragraph(therapist["bio"])
    doc.add_paragraph().add_run("Focus areas: ").bold = True
    doc.paragraphs[-1].add_run(therapist["specialties"])

    doc.add_heading("Demonstration credential", level=1)
    add_key_value_table(
        doc,
        [
            ("Credential type", therapist["credentials"]),
            ("License identifier", therapist["license_id"]),
            ("Jurisdiction", therapist["jurisdiction"]),
            ("Verification status", "Not verified - demonstration value only"),
            ("Connected demo client", therapist["client"]),
        ],
    )
    warning = doc.add_paragraph()
    warning.paragraph_format.space_before = Pt(5)
    warning.add_run("Important: ").bold = True
    warning.add_run(
        "The DEMO identifier is intentionally non-valid. It must never satisfy a real credential check or unlock therapist permissions."
    )

    doc.add_heading("Expected verification workflow", level=1)
    add_verification_steps(doc)
    doc.save(output)
    return output


def build_client_document(output_root: Path, client: dict) -> Path:
    output = output_root / "Clients" / client["folder"] / "client-intake-and-preferences.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    title = f"{client['name']} Client Intake and Preferences"
    configure_document(doc, title, "Fictional client onboarding and care-preference example")
    add_title_block(
        doc,
        title,
        "A fictional starting record for testing client onboarding, therapist connection, scheduling, and document workflows.",
    )
    doc.add_heading("Account and care relationship", level=1)
    add_key_value_table(
        doc,
        [
            ("Client name", client["name"]),
            ("Role", "Client"),
            ("Sign in email", client["email"]),
            ("Account state", "Fictional development fixture"),
            ("Auth0 subject", "Populate after development account creation"),
            ("Clinical fixture", "See the separate trajectory outline and session plan"),
        ],
    )

    doc.add_heading("Starting context", level=1)
    doc.add_paragraph(client["reason"])
    doc.add_heading("Initial goals", level=2)
    add_bullets(doc, client["goals"])

    doc.add_heading("Scheduling and communication preferences", level=1)
    add_key_value_table(doc, client["preferences"])

    doc.add_heading("Demo workflow notes", level=1)
    add_bullets(
        doc,
        [
            "Create the Auth0 identity with the sign in email, then copy the issued subject into the account fixture.",
            "Connect the client to a therapist only through the application's invitation or approval workflow.",
            "Keep passwords, tokens, diagnoses, session content, and signed forms out of Auth0 metadata and tokens.",
        ],
    )
    doc.save(output)
    return output


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-root", type=Path, required=True)
    args = parser.parse_args()
    outputs = []
    for therapist in THERAPISTS:
        outputs.append(build_therapist_document(args.output_root, therapist))
    for client in CLIENTS:
        outputs.append(build_client_document(args.output_root, client))
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
